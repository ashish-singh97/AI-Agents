"""
Excel (.xlsx) input reader. Reuses the same column-alias logic as the
CSV reader so both formats behave identically.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from database.models import RecruiterRecord
from input.csv_reader import _normalize_headers


def read_excel(path: str | Path, sheet_name: str | int = 0) -> List[RecruiterRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - import guard
        raise RuntimeError(
            "openpyxl is required to read Excel files. Run: pip install -r requirements.txt"
        ) from exc

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[sheet_name]] if isinstance(sheet_name, int) else wb[sheet_name]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValueError(f"Excel file has no rows: {path}")

    headers = [str(h) if h is not None else "" for h in header_row]
    header_map = _normalize_headers(headers)

    records: List[RecruiterRecord] = []
    for row in rows_iter:
        values = {"hr_name": "", "company": "", "job_role": "", "email": "", "linkedin_url": ""}
        for header, cell in zip(headers, row):
            canonical = header_map.get(header)
            if canonical and cell is not None:
                values[canonical] = str(cell).strip()

        if not any(values.values()):
            continue

        records.append(
            RecruiterRecord(
                hr_name=values["hr_name"],
                company=values["company"],
                job_role=values["job_role"],
                email=values["email"],
                linkedin_url=values["linkedin_url"],
            )
        )
    wb.close()
    return records
